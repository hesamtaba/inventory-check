import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill,Font,Alignment
from openpyxl.utils import get_column_letter
COLORS={'موجود':'C6EFCE','مغایرت':'FFC7CE','تکراری':'FCE4D6','شمارش‌نشده':'D9E1F2','کد نامعتبر':'FFC7CE'}
class ExportService:
 def __init__(self,db,projects,scans): self.db=db; self.projects=projects; self.scans=scans
 def export(self,pid,path):
  p=self.projects.get(pid); st=self.scans.stats(pid); wb=Workbook(); wb.remove(wb.active)
  ws=wb.create_sheet('خلاصه'); summary=[('نام پروژه',p['name']),('تاریخ ایجاد',p['created_at']),('نام فایل موجودی',Path(p['inventory_path'] or '').name),('تعداد کل موجودی',st['inventory']),('تعداد اسکن‌ها',st['scans']),('تعداد موجود',st['found']),('تعداد مغایرت',st['mismatch']),('تعداد تکراری',st['duplicate']),('تعداد شمارش‌نشده',st['unscanned']),('درصد پیشرفت',f"{st['progress']:.2f}%")]
  for r,x in enumerate(summary,1): ws.cell(r,1,x[0]); ws.cell(r,2,x[1])
  scans=self.db.query('SELECT * FROM scans WHERE project_id=? AND deleted=0 ORDER BY id',(pid,)); inv=self.db.query('SELECT * FROM inventory WHERE project_id=? ORDER BY row_no',(pid,))
  found_labels={x['normalized_label'] for x in scans if x['status'] in ('موجود','تکراری')}; un=[x for x in inv if x['normalized_label'] not in found_labels]
  def scan_sheet(name,rows):
   sh=wb.create_sheet(name); headers=['شماره لیبل واردشده','کد نرمال‌شده','وضعیت','تاریخ و ساعت','منبع ثبت','دفعات تکرار','توضیحات']; sh.append(headers)
   counts={}
   for x in scans: counts[x['normalized_label']]=counts.get(x['normalized_label'],0)+1
   for x in rows: sh.append([x['entered_label'],x['normalized_label'],x['status'],x['scanned_at'],x['source'],counts.get(x['normalized_label'],1),x['notes']])
   for row in sh.iter_rows(min_row=2):
    row[0].number_format=row[1].number_format='@'; fill=PatternFill('solid',fgColor=COLORS.get(row[2].value,'FFFFFF'))
    for c in row: c.fill=fill
   return sh
  scan_sheet('موجود',[x for x in scans if x['status']=='موجود']); scan_sheet('مغایرت',[x for x in scans if x['status'] in ('مغایرت','کد نامعتبر')]); scan_sheet('تکراری',[x for x in scans if x['status']=='تکراری'])
  sh=wb.create_sheet('شمارش‌نشده'); headers=['شماره لیبل اصلی','کد نرمال‌شده','ردیف فایل موجودی']; extra=[]
  if un: extra=list(json.loads(un[0]['row_json']).keys())
  sh.append(headers+extra)
  for x in un:
   d=json.loads(x['row_json']); sh.append([x['original_label'],x['normalized_label'],x['row_no']]+[d.get(k,'') for k in extra]); sh.cell(sh.max_row,1).number_format=sh.cell(sh.max_row,2).number_format='@'
   for c in sh[sh.max_row]: c.fill=PatternFill('solid',fgColor=COLORS['شمارش‌نشده'])
  scan_sheet('تمام اسکن‌ها',scans)
  for sh in wb.worksheets:
   sh.sheet_view.rightToLeft=True; sh.freeze_panes='A2'
   for c in sh[1]: c.font=Font(bold=True); c.alignment=Alignment(horizontal='center')
   for col in range(1,sh.max_column+1): sh.column_dimensions[get_column_letter(col)].width=min(max(12,max((len(str(sh.cell(r,col).value or '')) for r in range(1,min(sh.max_row,200)+1)),default=12)+2),35)
  wb.save(path); return path
